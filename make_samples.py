"""Generate sample data: 5 broker statement PDFs, 5 Arbor client-trade sheets, 5 MasterPro trackers.

Usage: python make_samples.py [root]   (default ./samples)
Planted deviations: SELL client < advisor and BUY client > advisor (suspicious), a missing trade and a quantity
difference (other discrepancies), plus favourable client prices and a clean account that must not be flagged.
"""
import os
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Font
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

root = sys.argv[1] if len(sys.argv) > 1 else "samples"
for sub in ("statements", "client_trades", "master"):
    os.makedirs(f"{root}/{sub}", exist_ok=True)

# trade: (symbol, iso date, side, qty, advisor price on statement, client price in Arbor, client qty or None=missing)
BROKERS = [
    dict(file="fidelity_oct2026", broker="Fidelity Investments", advisor="Daniel Ross", fmt="us", accounts=[
        dict(number="Z12-345678", type="Individual", holder="Margaret Chen", related=["James Chen (Joint Owner)"],
             trades=[("AAPL", "2026-10-05", "BUY", 100, 191.50, 191.50, 100),
                     ("MSFT", "2026-10-12", "SELL", 50, 405.00, 403.20, 50),
                     ("TSLA", "2026-10-20", "BUY", 20, 250.00, 250.00, 20)])]),
    dict(file="schwab_oct2026", broker="Charles Schwab & Co.", advisor="Priya Nair", fmt="long", accounts=[
        dict(number="5511-2290", type="Traditional IRA", holder="Robert Alvarez", related=["Elena Alvarez (Beneficiary)"],
             trades=[("NVDA", "2026-10-03", "BUY", 10, 900.00, 905.00, 10),
                     ("AMZN", "2026-10-14", "BUY", 30, 178.00, 177.50, 30),
                     ("GOOGL", "2026-10-21", "SELL", 15, 170.00, 170.00, 15)]),
        dict(number="5511-7781", type="Individual", holder="Robert Alvarez", related=[],
             trades=[("KO", "2026-10-09", "SELL", 200, 61.20, 61.20, 200)])]),
    dict(file="vanguard_oct2026", broker="Vanguard Brokerage Services", advisor="Helen Walsh", fmt="us", accounts=[
        dict(number="88123456", type="Trust", holder="Patel Family Trust",
             related=["Anita Patel (Trustee)", "Rahul Patel (Beneficiary)"],
             trades=[("VTI", "2026-10-06", "BUY", 50, 270.10, 270.10, 50),
                     ("VXUS", "2026-10-13", "BUY", 80, 62.40, 62.40, None),
                     ("BND", "2026-10-22", "SELL", 120, 72.00, 72.00, 100)])]),
    dict(file="ibkr_oct2026", broker="Interactive Brokers LLC", advisor="Tom Becker", fmt="dmy", accounts=[
        dict(number="U1234567", type="Individual", holder="Liam O'Connor", related=[],
             trades=[("META", "2026-10-08", "SELL", 25, 520.00, 521.00, 25),
                     ("AMD", "2026-10-16", "BUY", 60, 160.00, 159.50, 60)])]),
    dict(file="etrade_oct2026", broker="E*TRADE from Morgan Stanley", advisor="Carla Mendes", fmt="long", accounts=[
        dict(number="4821-9930", type="Joint", holder="Sofia Rinaldi", related=["Marco Rinaldi (Joint Owner)"],
             trades=[("NFLX", "2026-10-02", "BUY", 12, 640.00, 642.00, 12),
                     ("JPM", "2026-10-19", "SELL", 40, 195.00, 194.00, 40)])]),
]


def fmt_date(iso, style):
    d = date.fromisoformat(iso)
    return {"us": d.strftime("%m/%d/%Y"), "long": d.strftime("%b %d, %Y"), "dmy": d.strftime("%d-%b-%Y")}[style]


def write_pdf(b):
    c = canvas.Canvas(f"{root}/statements/{b['file']}.pdf", pagesize=letter)
    y = 750

    def line(text, size=10, bold=False, dy=16):
        nonlocal y
        if y < 80:
            c.showPage()
            y = 750
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(50, y, text)
        y -= dy

    line(b["broker"], 18, True, 28)
    line("Monthly Account Statement - October 2026", 11, dy=24)
    for a in b["accounts"]:
        line(f"Account Number: {a['number']}", 11, True)
        line(f"Account Type: {a['type']}")
        line(f"Account Holder: {a['holder']}")
        if a["related"]:
            line("Related Persons: " + "; ".join(a["related"]))
        line(f"Financial Advisor: {b['advisor']}", dy=22)
        line("TRADE ACTIVITY", 11, True)
        if b["fmt"] == "us":
            line("Trade Date    Action   Symbol   Quantity   Price", 9, True, 14)
            for s, d, side, q, ap, _, _ in a["trades"]:
                line(f"{fmt_date(d, 'us')}    {side.title():7}  {s:7}  {q:>8}   ${ap:,.2f}", 9, dy=14)
        elif b["fmt"] == "long":
            for s, d, side, q, ap, _, _ in a["trades"]:
                verb = "Bought" if side == "BUY" else "Sold"
                line(f"{fmt_date(d, 'long')} - {verb} {q} shares of {s} at {ap:,.2f} per share", 9, dy=14)
        else:
            line("Date | Symbol | Buy/Sell | Qty | Trade Price (USD)", 9, True, 14)
            for s, d, side, q, ap, _, _ in a["trades"]:
                line(f"{fmt_date(d, 'dmy')} | {s} | {'B' if side == 'BUY' else 'S'} | {q} | {ap:.2f}", 9, dy=14)
        line("Dividends and fees are not trades and are listed separately.", 8, dy=12)
        line("Total account value is available online.", 8, dy=26)
    c.save()


def write_client_trades(b):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Arbor Trades"
    ws.append(["Account Number", "Client Name", "Symbol", "Trade Date", "Side", "Quantity", "Price"])
    for a in b["accounts"]:
        for s, d, side, q, _, cp, cq in a["trades"]:
            if cq is not None:
                ws.append([a["number"], a["holder"], s, date.fromisoformat(d), side, cq, cp])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "yyyy-mm-dd"
    wb.save(f"{root}/client_trades/arbor_{b['file']}.xlsx")


def write_master(b):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MasterPro"
    ws.append(["Account Number", "Client Name", "Advisor", "Broker", "Account Type",
               "Review Date", "Number of Trades", "Suspicious Activity"])
    for a in b["accounts"]:
        ws.append([a["number"], a["holder"], b["advisor"], b["broker"], a["type"]])
    ws.append([f"ZZ-{b['accounts'][0]['number'][-4:]}", "Not On Statement", b["advisor"], b["broker"], "Individual"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(f"{root}/master/masterpro_{b['file'].split('_')[0]}.xlsx")


for b in BROKERS:
    write_pdf(b)
    write_client_trades(b)
    write_master(b)
print(f"samples written to {root}: {len(BROKERS)} PDFs, {len(BROKERS)} client trade sheets, {len(BROKERS)} trackers")
