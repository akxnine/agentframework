import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Discrepancy

COLS = ["kind", "account", "symbol", "trade_date", "side", "client_value", "advisor_value", "detail", "client_file", "advisor_file"]


def _sheet(ws, rows: list[Discrepancy]):
    ws.append([c.replace("_", " ").title() for c in COLS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append([getattr(r, c) for c in COLS])
    for i, c in enumerate(COLS, 1):
        w = max([len(str(x.value or "")) for x in ws[get_column_letter(i)]] + [10])
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, 60)
    ws.freeze_panes = "A2"


def build_excel(discrepancies: list[Discrepancy], out_dir: str, n_client: int, n_advisor: int) -> str:
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"suspicious_activity_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    summ.append(["Client trades", n_client])
    summ.append(["Advisor trades", n_advisor])
    summ.append(["Total discrepancies", len(discrepancies)])
    kinds = sorted({d.kind for d in discrepancies})
    for k in kinds:
        summ.append([k, sum(d.kind == k for d in discrepancies)])
    summ.column_dimensions["A"].width = 28
    _sheet(wb.create_sheet("Discrepancies"), [d for d in discrepancies if not d.kind.startswith("MISSING")])
    _sheet(wb.create_sheet("Unmatched"), [d for d in discrepancies if d.kind.startswith("MISSING")])
    wb.save(path)
    return path


def build_pdf(discrepancies: list[Discrepancy], out_dir: str, n_client: int, n_advisor: int) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"suspicious_activity_report_{datetime.now():%Y%m%d_%H%M%S}.pdf")
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=8, leading=10)

    def grid(data, header=True):
        t = Table(data, repeatRows=1 if header else 0)
        style = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                      ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
        t.setStyle(TableStyle(style))
        return t

    story = [Paragraph("Suspicious Activity Report", styles["Title"]),
             Paragraph(f"Generated {datetime.now():%Y-%m-%d %H:%M}", styles["Normal"]), Spacer(1, 12)]
    summary = [["Client trades", n_client], ["Advisor trades", n_advisor], ["Total discrepancies", len(discrepancies)]]
    summary += [[k, sum(d.kind == k for d in discrepancies)] for k in sorted({d.kind for d in discrepancies})]
    story += [grid(summary, header=False), Spacer(1, 16), Paragraph("Discrepancies", styles["Heading2"])]
    head = ["Kind", "Account", "Symbol", "Date", "Side", "Client", "Advisor", "Detail"]
    rows = [head] + [[Paragraph(str(x), cell) for x in (d.kind, d.account, d.symbol, d.trade_date, d.side, d.client_value,
                                                      d.advisor_value, d.detail)] for d in discrepancies]
    story.append(grid(rows))
    SimpleDocTemplate(path, pagesize=landscape(letter), title="Suspicious Activity Report").build(story)
    return path
