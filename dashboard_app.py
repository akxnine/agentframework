"""Interactive dashboard: streamlit run dashboard_app.py"""
import asyncio
import glob
import os
import time
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl import load_workbook

from orchestrator.config import Settings
from orchestrator.costs import DEFAULT_PRICING, PRICING, cost, scan_inputs
from orchestrator.events import read_events
from orchestrator.ledger import ComplianceLedger
from orchestrator.runner import RunController

st.set_page_config(page_title="Statement Reconciliation", page_icon="📊", layout="wide")


@st.cache_resource
def controller() -> RunController:
    return RunController()


ctl = controller()
OUT = Settings().output_folder
RUNS = os.path.join(OUT, "runs")


@st.cache_data(ttl=10)
def cached_scan(folder: str) -> dict:
    return scan_inputs(folder)


def run_files() -> list[str]:
    return sorted(glob.glob(os.path.join(RUNS, "*.jsonl")), reverse=True)


def summarize(evs: list[dict]) -> dict:
    """Fold the orchestrator's event stream into the state the UI renders."""
    s = {"transcript": [], "docs": [], "disc": [], "metrics": {}, "status": "idle", "result": "",
        "start": None, "end": None, "plan": None, "client_files": [], "trackers": [], "reports": None,
        "compliance_cases": [], "logs": []}
    for e in evs:
        k = e["kind"]
        if k == "run_start":
            s["status"], s["start"] = "running", e["ts"]
        elif k == "log":
            s["logs"].append(f"{time.strftime('%H:%M:%S', time.localtime(e['ts']))} {e['level']:<7} {e['msg']}")
        elif k == "agent_text":
            s["transcript"].append({"t": e["ts"], "type": "text", "text": e["text"]})
        elif k == "tool_use":
            s["transcript"].append({"t": e["ts"], "type": "tool_use", "name": e["name"], "input": e["input"]})
        elif k == "tool_result":
            s["transcript"].append({"t": e["ts"], "type": "tool_result", "is_error": e.get("is_error", False),
                                    "content": e["content"]})
        elif k == "doc":
            s["docs"].append(e)
        elif k == "discrepancies":
            s["disc"] = e["rows"]
        elif k == "metrics":
            s["metrics"] = e
        elif k == "plan":
            s["plan"] = e
        elif k == "client_file":
            s["client_files"].append(e)
        elif k == "tracker":
            s["trackers"].append(e)
        elif k == "reports":
            s["reports"] = e
        elif k == "compliance_case":
            s["compliance_cases"].append(e)
        elif k in ("agent_result", "run_end"):
            s["status"] = "failed" if k == "run_end" and e.get("status") == "failed" else (
                "cancelled" if k == "run_end" and e.get("status") == "cancelled" else "done")
            s["result"], s["end"] = e.get("result", s["result"]), e["ts"]
    return s


def file_link(path: str) -> str:
    return f"[{path}]({Path(path).resolve().as_uri()})"


def doc_frame(s: dict, pricing: str) -> pd.DataFrame:
    rows = []
    for i, d in enumerate(s["docs"], 1):
        rows.append({
            "#": i, "File": d["file"], "Status": "✅ OK" if d["status"] == "OK" else "❌ FAILED",
            "Statements": d.get("statements", 0), "Trades": d.get("trades", 0),
            "Tokens in": d.get("prompt_tokens", 0), "Tokens out": d.get("completion_tokens", 0),
            f"Cost ($) [{pricing}]": round(cost(d.get("prompt_tokens", 0), d.get("completion_tokens", 0), pricing), 4),
            "Notes": d.get("error") or ", ".join(d.get("accounts", [])),
        })
    return pd.DataFrame(rows)


def prerun_panel(scan: dict, title: str) -> None:
    docs = scan["docs"]
    pages = sum(d["pages"] or 0 for d in docs)
    with st.container(border=True):
        st.markdown(f"**{title}**")
        st.info(f"📄 **{len(docs)}** PDF(s) to scan ({pages} pages) · 📑 {len(scan['client_files'])} client trade "
                f"file(s) · 🗂 {len(scan['master_files'])} master tracker(s)")
        c1, c2 = st.columns([3, 2])
        c1.dataframe(pd.DataFrame([{"Check": n, "Status": "✅" if ok else "⚠️", "Detail": d}
                                   for n, ok, d in scan["checks"]]), hide_index=True, width="stretch")
        if docs:
            c2.dataframe(pd.DataFrame(docs), hide_index=True, width="stretch")


def download_buttons(paths: list[str], key: str) -> None:
    cols = st.columns([1, 1, 6])
    for col, path in zip(cols, paths):
        if os.path.exists(path):
            with open(path, "rb") as fh:
                col.download_button(f"⬇ {os.path.splitext(path)[1][1:].upper()}", fh.read(),
                                    os.path.basename(path), key=f"{key}_{path}")


def result_row(s: dict, run_id: str) -> None:
    if s["status"] == "failed":
        st.error(s["result"])
        return
    with st.container(border=True):
        st.markdown("**Result:** " + s["result"])
        rep = s["reports"]
        if rep:
            st.markdown(f"**Excel:** {file_link(rep['excel'])}  \n**PDF:** {file_link(rep['pdf'])}")
            download_buttons([rep["excel"], rep["pdf"]], f"res_{run_id}")


def transcript_view(s: dict) -> None:
    st.markdown("##### Orchestrator transcript")
    if not s["transcript"]:
        st.caption("No activity yet.")
        return
    for item in s["transcript"][-150:]:
        ts = time.strftime("%H:%M:%S", time.localtime(item["t"]))
        if item["type"] == "text":
            st.markdown(f"`{ts}` 🧠 {item['text']}")
        elif item["type"] == "tool_use":
            st.markdown(f"`{ts}` 🔧 **{item['name']}**`({item['input']})`")
        elif item["type"] == "tool_result":
            icon = "❌" if item["is_error"] else "✅"
            st.markdown(f"`{ts}` {icon} _{item['content']}_")


def live_view(path: str | None, was_running: bool, pricing: str, folder: str) -> None:
    s = summarize(read_events(path)) if path else summarize([])
    badge = {"idle": "⚪ Idle", "running": "🟢 Running", "done": "✅ Finished", "failed": "🔴 Failed",
             "cancelled": "🟠 Cancelled"}[s["status"]]
    elapsed = ((s["end"] or time.time()) - s["start"]) if s["start"] else 0
    run_id = os.path.basename(path)[:-6] if path else ""
    st.subheader(f"{badge}  ·  {elapsed:.0f}s" + (f"  ·  run {run_id}" if path else ""))

    if s["plan"]:
        prerun_panel(s["plan"], "Pre-run summary (captured at run start)")
    elif not ctl.running:
        prerun_panel(cached_scan(folder), "Pre-run check (current docs folder)")

    m = s["metrics"]
    if m:
        for col, (label, key) in zip(st.columns(4), [("Advisor trades", "advisor_trades"),
                                                     ("Client trades", "client_trades"), ("Discrepancies", "discrepancies"),
                                                     ("Suspicious", "suspicious")]):
            col.metric(label, m.get(key, 0))

    transcript_view(s)

    df = doc_frame(s, pricing)
    if not df.empty:
        st.markdown("##### Statement extraction")
        st.dataframe(df, hide_index=True, width="stretch")
    if s["client_files"]:
        st.markdown("##### Client trades loaded")
        st.dataframe(pd.DataFrame(s["client_files"]).drop(columns=["kind", "ts"]), hide_index=True, width="stretch")
    if s["disc"]:
        st.markdown("##### Discrepancies")
        st.dataframe(pd.DataFrame(s["disc"]), hide_index=True, width="stretch")
    if s["trackers"]:
        st.markdown("##### MasterPro trackers updated")
        st.dataframe(pd.DataFrame(s["trackers"]).drop(columns=["kind", "ts"]), hide_index=True, width="stretch")
    if s["compliance_cases"]:
        st.markdown("##### Compliance cases opened this run")
        st.dataframe(pd.DataFrame(s["compliance_cases"]).drop(columns=["kind", "ts"]), hide_index=True, width="stretch")

    if s["result"]:
        result_row(s, run_id)

    with st.expander(f"Raw log ({len(s['logs'])} lines)", expanded=False):
        st.code("\n".join(s["logs"][-200:]) or "No log yet", language="text", height=220)

    if was_running and not ctl.running:
        st.rerun()  # run finished: refresh the other tabs and the sidebar


def live_tab(path: str | None, pricing: str, folder: str) -> None:
    was_running = ctl.running
    st.fragment(run_every=1 if was_running else None)(lambda: live_view(path, was_running, pricing, folder))()


def metrics_tab(s: dict, pricing: str) -> None:
    if not s["docs"]:
        st.info("No extraction metrics for this run yet.")
        return
    df = doc_frame(s, pricing)
    docs = pd.DataFrame(s["docs"])
    pin = int(docs.get("prompt_tokens", pd.Series(dtype=int)).sum())
    pout = int(docs.get("completion_tokens", pd.Series(dtype=int)).sum())
    st.caption(f"Extraction model: **{Settings().model}** (Claude, via the Agent SDK orchestrator). "
              f"Cost columns use the pricing selected in the sidebar: **{pricing}**.")
    a, b, c = st.columns(3)
    a.metric("PDFs OK", f"{(docs.status == 'OK').sum()}/{len(docs)}")
    b.metric("Tokens in / out", f"{pin:,} / {pout:,}")
    c.metric(f"Est. cost ({pricing})", f"${cost(pin, pout, pricing):.4f}")
    st.dataframe(df, hide_index=True, width="stretch")
    st.markdown("##### Cost if this run used a different model")
    st.dataframe(pd.DataFrame([{"Model": n, "$/1M in": pr[0], "$/1M out": pr[1],
                                "Est. cost for this run ($)": round(cost(pin, pout, n), 4)}
                               for n, pr in PRICING.items()]), hide_index=True, width="stretch")
    st.caption("Prices are approximate list prices (edit `orchestrator/costs.py`).")
    if "prompt_tokens" in docs:
        st.plotly_chart(px.bar(docs, x="file", y=["prompt_tokens", "completion_tokens"], title="Tokens per PDF"),
                        width="stretch")


def discrepancies_tab(s: dict) -> None:
    if not s["disc"]:
        st.info("No discrepancies recorded for this run.")
        return
    df = pd.DataFrame(s["disc"])
    f1, f2 = st.columns(2)
    kinds = f1.multiselect("Kind", sorted(df.kind.unique()), default=sorted(df.kind.unique()))
    accts = f2.multiselect("Account", sorted(df.account.unique()), default=sorted(df.account.unique()))
    view = df[df.kind.isin(kinds) & df.account.isin(accts)]
    st.dataframe(view.style.apply(lambda r: ["background-color:#5c1f1f" if r.kind == "PRICE" else "" for _ in r], axis=1),
                 width="stretch", hide_index=True)
    st.plotly_chart(px.histogram(view, x="kind", color="account", title="Discrepancies by kind"), width="stretch")
    st.download_button("Download CSV", view.to_csv(index=False), "discrepancies.csv")


@st.cache_data
def read_report(path: str, mtime: float) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    summary = {str(r[0]): r[1] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0] is not None}

    def sheet(name: str) -> pd.DataFrame:
        if name not in wb.sheetnames:
            return pd.DataFrame()
        rows = list(wb[name].iter_rows(values_only=True))
        return pd.DataFrame(rows[1:], columns=rows[0]) if rows else pd.DataFrame()

    return {"summary": summary, "discrepancies": sheet("Discrepancies"), "unmatched": sheet("Unmatched"),
            "notes": sheet("Salesforce Notes")}


def reports_tab() -> None:
    stems = sorted({os.path.splitext(f)[0] for f in glob.glob(os.path.join(OUT, "suspicious_activity_report_*"))},
                   reverse=True)
    if not stems:
        st.info("No reports generated yet.")
        return
    totals = ("Client trades", "Advisor trades", "Total discrepancies")
    rows = []
    for stem in stems:
        xlsx, pdf = stem + ".xlsx", stem + ".pdf"
        info = read_report(xlsx, os.path.getmtime(xlsx))["summary"] if os.path.exists(xlsx) else {}
        kinds = {k: v for k, v in info.items() if k not in totals}
        ts = os.path.basename(stem).rsplit("_", 2)
        rows.append({
            "Generated": pd.to_datetime(f"{ts[-2]}{ts[-1]}", format="%Y%m%d%H%M%S", errors="coerce"),
            "Client trades": info.get("Client trades"), "Advisor trades": info.get("Advisor trades"),
            "Discrepancies": info.get("Total discrepancies"), "Suspicious (PRICE)": kinds.get("PRICE", 0),
            "By kind": ", ".join(f"{k}: {v}" for k, v in kinds.items()),
            "Excel (KB)": round(os.path.getsize(xlsx) / 1024, 1) if os.path.exists(xlsx) else None,
            "PDF (KB)": round(os.path.getsize(pdf) / 1024, 1) if os.path.exists(pdf) else None,
            "Excel": os.path.abspath(xlsx), "PDF": os.path.abspath(pdf),
        })
    st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch")
    pick = st.selectbox("Open report", stems,
                        format_func=lambda x: os.path.basename(x).replace("suspicious_activity_report_", ""))
    xlsx, pdf = pick + ".xlsx", pick + ".pdf"
    st.markdown(f"**Excel:** {file_link(xlsx)}  \n**PDF:** {file_link(pdf)}")
    download_buttons([xlsx, pdf], "rep")
    if os.path.exists(xlsx):
        data = read_report(xlsx, os.path.getmtime(xlsx))
        st.markdown("##### Discrepancies")
        st.dataframe(data["discrepancies"], hide_index=True, width="stretch")
        st.markdown("##### Unmatched trades")
        st.dataframe(data["unmatched"], hide_index=True, width="stretch")
        if not data["notes"].empty:
            st.markdown("##### Salesforce notes")
            st.dataframe(data["notes"], hide_index=True, width="stretch")


def history_tab(pricing: str) -> None:
    rows = []
    for f in run_files():
        s = summarize(read_events(f))
        pin = sum(d.get("prompt_tokens", 0) for d in s["docs"])
        pout = sum(d.get("completion_tokens", 0) for d in s["docs"])
        rows.append({
            "Run": os.path.basename(f)[:-6],
            "Started": pd.to_datetime(s["start"], unit="s") if s["start"] else None,
            "Status": s["status"],
            "Duration (s)": round(s["end"] - s["start"], 1) if s["end"] and s["start"] else None,
            "PDFs": len(s["docs"]), "Tokens": pin + pout, f"Cost ($) [{pricing}]": round(cost(pin, pout, pricing), 4),
            "Discrepancies": s["metrics"].get("discrepancies"), "Suspicious": s["metrics"].get("suspicious"),
        })
    if rows:
        st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch")
    else:
        st.info("No runs yet.")


def compliance_tab(settings: Settings) -> None:
    ledger = ComplianceLedger(settings.ledger_path)
    st.caption("Every flagged account emailed to compliance, and the decision that closed it out. "
              "The watcher process (`python cli.py watch-compliance`) polls the mailbox and fills this "
              "in automatically; use the button below to check once from here instead.")
    c1, c2 = st.columns([1, 4])
    if c1.button("📥 Check mailbox for replies now", disabled=bool(settings.require_email_env())):
        from orchestrator.compliance_watch import process_once
        from anthropic import AsyncAnthropic
        client = AsyncAnthropic(api_key=settings.anthropic_api_key) if settings.anthropic_api_key else None
        with st.spinner("Polling mailbox..."):
            n = asyncio.run(process_once(settings, ledger, client))
        st.success(f"Processed {n} reply(ies).") if n else st.info("No new replies.")
    cases = ledger.list_cases()
    if not cases:
        st.info("No compliance cases yet. They're created automatically when a run finds suspicious pricing.")
        return
    df = pd.DataFrame([c.model_dump() for c in cases])
    df["report_paths"] = df["report_paths"].apply(lambda p: ", ".join(os.path.basename(x) for x in p))
    status_filter = st.multiselect("Status", sorted(df.status.unique()), default=sorted(df.status.unique()))
    st.dataframe(df[df.status.isin(status_filter)], hide_index=True, width="stretch")
    pending = (df.status == "PENDING").sum()
    if pending:
        st.warning(f"{pending} case(s) awaiting compliance review.")


with st.sidebar:
    st.title("📊 Reconciliation")
    folder = st.text_input("Docs folder", Settings().docs_folder)
    dry = st.toggle("Dry run (no email)", True)
    in_place = st.toggle("Update master trackers in place", False)
    pricing = st.selectbox("Cost model", list(PRICING), index=list(PRICING).index(DEFAULT_PRICING))
    scan = cached_scan(folder)
    if not ctl.running:
        if st.button("▶ Run reconciliation", type="primary", width="stretch", disabled=not scan["blocking_ok"]):
            cfg = Settings()
            cfg.docs_folder, cfg.in_place = folder, in_place
            ctl.start(cfg, dry)
            st.rerun()
        if not scan["blocking_ok"]:
            st.caption("⚠️ Need statements, client trades, master trackers, and an Anthropic API key.")
    else:
        st.button("⏳ Running…", disabled=True, width="stretch")
        if st.button("⏹ Cancel", width="stretch"):
            ctl.cancel()
    st.divider()
    files = run_files()
    live_id = os.path.join(RUNS, f"{ctl.run_id}.jsonl") if ctl.running else None
    choice = st.selectbox("Run to view", files, format_func=lambda f: os.path.basename(f)[:-6]) if files else None
    selected = live_id or choice

tabs = st.tabs(["🔴 Live", "📈 Metrics", "⚠️ Discrepancies", "📄 Reports", "🕘 History", "✅ Compliance"])
state = summarize(read_events(selected)) if selected else summarize([])
with tabs[0]:
    live_tab(selected, pricing, folder)
with tabs[1]:
    metrics_tab(state, pricing)
with tabs[2]:
    discrepancies_tab(state)
with tabs[3]:
    reports_tab()
with tabs[4]:
    history_tab(pricing)
with tabs[5]:
    compliance_tab(Settings())
