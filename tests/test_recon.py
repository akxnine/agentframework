import openpyxl
import pymupdf

from recon.compare import compare_trades
from recon.models import Trade
from recon.report import build_excel, build_pdf


def T(sym, side, qty, px, acct="A1", d="2026-01-02"):
    return Trade(symbol=sym, trade_date=d, side=side, quantity=qty, price=px, account=acct)


def kinds(c, a):
    return sorted(d.kind for d in compare_trades(c, a))


def test_buy_client_pays_more_flagged():
    assert kinds([T("X", "BUY", 1, 10.5)], [T("X", "BUY", 1, 10)]) == ["PRICE"]


def test_buy_client_pays_less_ok():
    assert kinds([T("X", "BUY", 1, 9.5)], [T("X", "BUY", 1, 10)]) == []


def test_sell_client_gets_less_flagged():
    assert kinds([T("X", "SELL", 1, 9.5)], [T("X", "SELL", 1, 10)]) == ["PRICE"]


def test_sell_client_gets_more_ok():
    assert kinds([T("X", "SELL", 1, 10.5)], [T("X", "SELL", 1, 10)]) == []


def test_qty_account_and_missing():
    c = [T("X", "BUY", 2, 10, "A1"), T("Y", "BUY", 1, 1)]
    a = [T("X", "BUY", 3, 10, "A2"), T("Z", "SELL", 1, 1)]
    assert kinds(c, a) == ["ACCOUNT", "MISSING_IN_ADVISOR", "MISSING_IN_CLIENT", "QUANTITY"]


def test_excel(tmp_path):
    d = compare_trades([T("X", "BUY", 1, 11)], [T("X", "BUY", 1, 10)])
    p = build_excel(d, str(tmp_path), 1, 1)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Summary", "Discrepancies", "Unmatched"]
    assert wb["Discrepancies"].max_row == 2


def test_read_pdf_text_and_scanned(tmp_path):
    import pytest
    from recon.extraction import read_pdf_text

    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "Trade AAPL")
    p = tmp_path / "a.pdf"
    doc.save(p)
    assert "AAPL" in read_pdf_text(str(p))
    blank = pymupdf.open()
    blank.new_page()
    q = tmp_path / "b.pdf"
    blank.save(q)
    with pytest.raises(ValueError, match="scanned"):
        read_pdf_text(str(q))


def test_compare_by_account_and_tracker(tmp_path):
    from recon.excel_io import update_tracker
    from recon.pipeline import compare_by_account

    c = [T("X", "BUY", 1, 11, "A-1"), T("X", "BUY", 1, 10, "B-2")]
    a = [T("X", "BUY", 1, 10, "A 1"), T("X", "BUY", 1, 10, "B-2")]
    d = compare_by_account(c, a)
    assert [(x.kind, x.account) for x in d] == [("PRICE", "A 1")]

    wb = openpyxl.Workbook()
    wb.active.append(["Account Number", "Client Name"])
    wb.active.append(["A-1", "n"])
    wb.active.append(["Z-9", "m"])
    src = tmp_path / "m.xlsx"
    wb.save(src)
    out = tmp_path / "out" / "m.xlsx"
    assert update_tracker(str(src), str(out), {"A1": (3, True)}, "2026-11-01") == ["A-1"]
    ws = openpyxl.load_workbook(out).active
    assert [c.value for c in ws[2]] == ["A-1", "n", "2026-11-01", 3, "Yes"]
    assert [c.value for c in ws[3]][2:] == [None, None, None]


def test_pdf(tmp_path):
    d = compare_trades([T("X", "BUY", 1, 11)], [T("X", "BUY", 1, 10)])
    p = build_pdf(d, str(tmp_path), 1, 1)
    with pymupdf.open(p) as doc:
        assert "PRICE" in doc[0].get_text()
