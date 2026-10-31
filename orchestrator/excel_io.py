import os
import shutil
from copy import copy
from datetime import date, datetime

import openpyxl

from .models import Trade

CLIENT_COLS = ["Account Number", "Client Name", "Symbol", "Trade Date", "Side", "Quantity", "Price"]
OUT_COLS = ["Review Date", "Number of Trades", "Suspicious Activity"]
COMPLIANCE_COLS = ["Compliance Status", "Reviewed By", "Reviewed At", "Compliance Notes"]


def norm_acct(a) -> str:
    return "".join(ch for ch in str(a) if ch.isalnum()).upper()


def _headers(ws) -> dict[str, int]:
    return {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}


def _iso(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def read_client_trades(path: str) -> list[Trade]:
    ws = openpyxl.load_workbook(path, data_only=True).active
    h = _headers(ws)
    missing = [c for c in CLIENT_COLS if c.lower() not in h]
    if missing:
        raise ValueError(f"{os.path.basename(path)}: missing columns {missing}")
    trades = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[h["account number"] - 1] in (None, ""):
            continue
        side = str(row[h["side"] - 1]).strip().upper()
        trades.append(Trade(
            symbol=str(row[h["symbol"] - 1]).strip().upper(), trade_date=_iso(row[h["trade date"] - 1]),
            side={"B": "BUY", "S": "SELL"}.get(side, side), quantity=float(row[h["quantity"] - 1]),
            price=float(row[h["price"] - 1]), account=str(row[h["account number"] - 1]).strip(),
            source_file=os.path.basename(path)))
    return trades


def _ensure_columns(ws, h: dict[str, int], names: list[str]) -> None:
    for name in names:
        if name.lower() not in h:
            h[name.lower()] = ws.max_column + 1
            ws.cell(row=1, column=h[name.lower()], value=name).font = copy(ws.cell(row=1, column=1).font)


def update_tracker(path: str, out_path: str, results: dict[str, tuple[int, bool]], review_date: str) -> list[str]:
    """Write Review Date / Number of Trades / Suspicious Activity for accounts in `results`
    (normalised account -> (trade count, suspicious)). Returns the accounts updated."""
    if os.path.abspath(path) != os.path.abspath(out_path):
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        shutil.copyfile(path, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    h = _headers(ws)
    if "account number" not in h:
        raise ValueError(f"{os.path.basename(path)}: no 'Account Number' column")
    _ensure_columns(ws, h, OUT_COLS)
    updated = []
    for row in range(2, ws.max_row + 1):
        acct = norm_acct(ws.cell(row=row, column=h["account number"]).value or "")
        if acct in results:
            count, suspicious = results[acct]
            ws.cell(row=row, column=h["review date"], value=review_date)
            ws.cell(row=row, column=h["number of trades"], value=count)
            ws.cell(row=row, column=h["suspicious activity"], value="Yes" if suspicious else "No")
            updated.append(ws.cell(row=row, column=h["account number"]).value)
    wb.save(out_path)
    return updated


def update_compliance_status(path: str, account: str, status: str, reviewer: str, reviewed_at: str,
                             notes: str = "") -> bool:
    """Write the compliance decision for one account's row into a tracker workbook in place.
    Used by the compliance watcher after a decision is parsed from an officer's email reply.
    Returns True if a matching row was found and updated."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    h = _headers(ws)
    if "account number" not in h:
        raise ValueError(f"{os.path.basename(path)}: no 'Account Number' column")
    _ensure_columns(ws, h, COMPLIANCE_COLS)
    target = norm_acct(account)
    found = False
    for row in range(2, ws.max_row + 1):
        if norm_acct(ws.cell(row=row, column=h["account number"]).value or "") == target:
            ws.cell(row=row, column=h["compliance status"], value=status)
            ws.cell(row=row, column=h["reviewed by"], value=reviewer)
            ws.cell(row=row, column=h["reviewed at"], value=reviewed_at)
            ws.cell(row=row, column=h["compliance notes"], value=notes)
            found = True
    if found:
        wb.save(path)
    return found


def tracker_accounts(path: str) -> set[str]:
    ws = openpyxl.load_workbook(path, data_only=True).active
    h = _headers(ws)
    col = h.get("account number")
    if not col:
        return set()
    return {norm_acct(r[col - 1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[col - 1]}
